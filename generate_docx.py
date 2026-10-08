from docx import Document
from docx.shared import Pt
import datetime

doc = Document()

# Title
title = doc.add_heading('Project Prototype Documentation', 0)
title.alignment = 1 # Center

# Subtitle
subtitle = doc.add_paragraph('Problem Statement 26250: "Air Power - Dynamic Air Operations & Resource Optimisation"\nMinistry of Defence | SIH 2026')
subtitle.alignment = 1

doc.add_paragraph(f'Date: {datetime.datetime.now().strftime("%Y-%m-%d")}')

doc.add_heading('1. Executive Summary', level=1)
doc.add_paragraph('This document outlines the current state of the prototype for an AI-enabled decision-support system for planning and dynamically retasking air operations in a contested, fast-changing environment.')

doc.add_heading('2. System Architecture & Tech Stack', level=1)
doc.add_paragraph('The prototype builds upon an existing predictive-maintenance and fleet-availability app. The architecture has been extended to support mission planning and dynamic re-allocation.')
ul1 = doc.add_paragraph(style='List Bullet')
ul1.add_run('Backend: ').bold = True
ul1.add_run('FastAPI, SQLAlchemy (SQLite for development)')
ul2 = doc.add_paragraph(style='List Bullet')
ul2.add_run('Optimization Engine: ').bold = True
ul2.add_run('PuLP with CBC solver (Mixed-Integer Linear Programming)')
ul3 = doc.add_paragraph(style='List Bullet')
ul3.add_run('Authentication: ').bold = True
ul3.add_run('OAuth2 with JWT (Analyst and Viewer roles)')

doc.add_heading('3. Completed Phases', level=1)

doc.add_heading('Phase 1: Operational Database & Data Seeding', level=2)
doc.add_paragraph('Designed and implemented a rich set of operational models to represent the airspace environment and resources:')
p1_ul = doc.add_paragraph(style='List Bullet')
p1_ul.add_run('Models created: ').bold = True
p1_ul.add_run('AirBase, Crew, WeaponType, Mission, AirspaceZone, ThreatZone, WeatherCondition, MissionPlan, PlanAssignment, DataFeed, EventLog.')
p1_ul2 = doc.add_paragraph(style='List Bullet')
p1_ul2.add_run('Aircraft Readiness: ').bold = True
p1_ul2.add_run('Aircraft are now evaluated for readiness (CAPABLE, LIMITED, GROUNDED) based on Remaining Useful Life (RUL) and failure risk.')
p1_ul3 = doc.add_paragraph(style='List Bullet')
p1_ul3.add_run('Synthetic Data: ').bold = True
p1_ul3.add_run('Implemented seed_ops.py to generate realistic synthetic scenarios for demonstration purposes.')

doc.add_heading('Phase 2: Optimization Engine & Planning API', level=2)
doc.add_paragraph('Developed a Mixed-Integer Linear Programming (MILP) solver to automate mission planning based on resource constraints and multiple objectives.')
p2_ul1 = doc.add_paragraph(style='List Bullet')
p2_ul1.add_run('Objectives: ').bold = True
p2_ul1.add_run('The engine generates three distinct plans per run: MAX_COVERAGE, MIN_RISK, and RESOURCE_LEAN.')
p2_ul2 = doc.add_paragraph(style='List Bullet')
p2_ul2.add_run('Constraints: ').bold = True
p2_ul2.add_run('Enforces hard constraints such as aircraft RUL margins, weapon compatibility, weather conditions (GO/NO-GO), and crew availability.')
p2_ul3 = doc.add_paragraph(style='List Bullet')
p2_ul3.add_run('API Endpoints: ').bold = True
p2_ul3.add_run('Exposed endpoints to generate plans, compare them side-by-side, and approve/reject proposed plans.')
p2_ul4 = doc.add_paragraph(style='List Bullet')
p2_ul4.add_run('Performance: ').bold = True
p2_ul4.add_run('Optimized for speed, generating all three plans in under 1.5 seconds for the demonstration scenario.')

doc.add_heading('Phase 3: Dynamic Retasking & Event Injection', level=2)
doc.add_paragraph('Implemented the ability to inject operational events and dynamically re-plan affected missions without discarding unaffected assignments.')
p3_ul1 = doc.add_paragraph(style='List Bullet')
p3_ul1.add_run('Supported Events: ').bold = True
p3_ul1.add_run('Aircraft Unserviceable, New Threat, Weather Closure, Airspace Closure, New Priority Mission, Crew Unavailable, and Feed Degraded.')
p3_ul2 = doc.add_paragraph(style='List Bullet')
p3_ul2.add_run('Incremental Optimization: ').bold = True
p3_ul2.add_run('Only affected missions are replanned, while assignments for unaffected missions remain locked to preserve continuity.')
p3_ul3 = doc.add_paragraph(style='List Bullet')
p3_ul3.add_run('Plan Versioning: ').bold = True
p3_ul3.add_run('Plans are versioned. Superseded plans are marked accordingly, and full diffs (ADDED, REMOVED, REPLACED, RETAINED) are computed and exposed via API.')

doc.add_heading('Phase 4: Common Operating Picture (COP) UI', level=2)
doc.add_paragraph('Completely redesigned the frontend into three tactical views:')
p4_ul1 = doc.add_paragraph(style='List Bullet')
p4_ul1.add_run('PICTURE: ').bold = True
p4_ul1.add_run('Real-time situational awareness dashboard.')
p4_ul2 = doc.add_paragraph(style='List Bullet')
p4_ul2.add_run('PLAN: ').bold = True
p4_ul2.add_run('Mission planning and approval interface.')
p4_ul3 = doc.add_paragraph(style='List Bullet')
p4_ul3.add_run('READINESS: ').bold = True
p4_ul3.add_run('Fleet availability and maintenance metrics.')

doc.add_heading('Phase 5: Data Fusion Layer', level=2)
doc.add_paragraph('Implemented a data fusion layer simulating incoming feeds from various sources.')
p5_ul1 = doc.add_paragraph(style='List Bullet')
p5_ul1.add_run('Feed Normalization: ').bold = True
p5_ul1.add_run('Normalizes diverse payloads (Aircraft Health, Threat Intel, Weather Intel) into actionable internal events.')
p5_ul2 = doc.add_paragraph(style='List Bullet')
p5_ul2.add_run('Deduplication: ').bold = True
p5_ul2.add_run('Hashes incoming payloads to prevent redundant event processing.')

doc.add_heading('Phase 6: Security Hardening', level=2)
doc.add_paragraph('Upgraded the security architecture to enterprise/defence-grade standards.')
p6_ul1 = doc.add_paragraph(style='List Bullet')
p6_ul1.add_run('Authentication: ').bold = True
p6_ul1.add_run('Implemented Argon2id for password hashing and asymmetric RS256 JWTs.')
p6_ul2 = doc.add_paragraph(style='List Bullet')
p6_ul2.add_run('Authorization: ').bold = True
p6_ul2.add_run('Enforced ABAC (Attribute-Based Access Control) via clearance levels and strict four-eyes verification for plan approval.')
p6_ul3 = doc.add_paragraph(style='List Bullet')
p6_ul3.add_run('Audit Trails: ').bold = True
p6_ul3.add_run('Created a tamper-evident, cryptographic (SHA-256) audit log chain.')

doc.add_heading('Phase 7: Testing & Documentation', level=2)
doc.add_paragraph('Finalized the PyTest suite covering the entire API (auth, planning, events, fusion, security) and compiled this comprehensive project documentation.')

doc.add_page_break()
doc.add_heading('Disclaimer', level=1)
doc.add_paragraph('All data contained within this prototype is synthetic. This system is not connected to operational defence systems and is intended solely for demonstration and training purposes.')

doc.save('/Users/shreyas/Desktop/26250/Prototype_Documentation.docx')
print("Documentation generated successfully.")
